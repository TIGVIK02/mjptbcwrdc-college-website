#!/usr/bin/env python3
"""Create a public-safe placement dataset from the internal Excel workbook."""

from __future__ import annotations

import json
import re
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SOURCE_XLSX = ROOT / "website raw data" / "Student_Placement.xlsx"
OUTPUT_JSON = ROOT / "data" / "student-placements.json"

SENSITIVE_PATTERNS = (
    "phone",
    "mobile",
    "contact",
    "email",
    "gmail",
    "address",
    "dob",
    "date of birth",
    "gender",
    "district",
    "aadhar",
    "aadhaar",
    "pan",
    "password",
    "social",
    "residential",
    "id number",
)

KEYWORDS = (
    "student",
    "name",
    "qualification",
    "company",
    "role",
    "package",
    "passing",
    "year",
)


def clean(value):
    if value is None:
        return ""
    if hasattr(value, "isoformat"):
        return value.isoformat()[:10]
    text = str(value).replace("\xa0", " ").replace("\n", " ").replace("\r", " ").strip()
    return re.sub(r"\s+", " ", text)


def normalize_key(key):
    return re.sub(r"[^a-z0-9]+", " ", clean(key).lower()).strip()


def matches_sensitive(header_label):
    key = normalize_key(header_label)
    return any(token in key for token in SENSITIVE_PATTERNS)


def first_present(*values):
    for value in values:
        text = clean(value)
        if text:
            return text
    return ""


def combine_name(parts):
    values = [clean(part) for part in parts if clean(part)]
    return " ".join(values).strip()


def normalize_record(raw_record, fallback_company):
    fields = {}
    for raw_key, value in raw_record.items():
        key = normalize_key(raw_key)
        if not key:
            continue
        if matches_sensitive(raw_key):
            continue
        fields[key] = clean(value)

    student = first_present(
        fields.get("student name"),
        fields.get("name of the student"),
        fields.get("student"),
        fields.get("name"),
        combine_name((fields.get("first name"), fields.get("last name"))),
    )

    company = first_present(
        fields.get("company name"),
        fields.get("company"),
        fields.get("organization"),
        fallback_company,
    )

    course = first_present(
        fields.get("qualification"),
        fields.get("course"),
        fields.get("qualification example bcom commerce btech eee bsc science etc"),
    )

    role = first_present(
        fields.get("job role"),
        fields.get("role"),
        fields.get("designation"),
    )

    package = first_present(
        fields.get("package"),
        fields.get("salary"),
        fields.get("annual package"),
        fields.get("ctc"),
    )

    academic_year = first_present(
        fields.get("year of passing"),
        fields.get("year of passing "),
        fields.get("academic year"),
        fields.get("pass year"),
    )

    if not student or not company:
        return None

    record = {
        "student": student,
        "company": company,
        "course": course,
        "role": role,
        "package": package,
        "academicYear": academic_year,
    }

    return {key: value for key, value in record.items() if value}


def parse_workbook(path: Path):
    workbook = load_workbook(path, data_only=True)
    rows = []

    for sheet in workbook.sheetnames:
        worksheet = workbook[sheet]
        sheet_rows = list(worksheet.iter_rows(values_only=True))
        if not sheet_rows:
            continue

        header_index = None
        for index, row in enumerate(sheet_rows):
            if any(
                cell is not None and isinstance(cell, str) and any(keyword in clean(cell).lower() for keyword in KEYWORDS)
                for cell in row
            ):
                header_index = index
                break

        if header_index is None:
            continue

        headers = [clean(cell) for cell in sheet_rows[header_index]]
        if not headers:
            continue

        previous_company = ""
        for raw_row in sheet_rows[header_index + 1 :]:
            if not raw_row or all(clean(cell) == "" for cell in raw_row):
                continue

            mapped = {}
            row_company = previous_company
            for idx, header in enumerate(headers):
                if idx >= len(raw_row):
                    continue
                if header == "":
                    continue

                if sheet == "Sheet3" and normalize_key(header) in {"company name", "company", "organization"}:
                    value = raw_row[idx]
                    if value is not None and clean(value) != "":
                        row_company = clean(value)
                        mapped["Company name"] = row_company
                    continue

                if matches_sensitive(header):
                    continue

                value = raw_row[idx]
                if value is not None and clean(value) != "":
                    mapped[header] = value

            if sheet == "Sheet3" and not row_company:
                row_company = previous_company
            if sheet == "Sheet3" and row_company:
                mapped["Company name"] = row_company
                previous_company = row_company

            record = normalize_record(mapped, row_company if sheet == "Sheet3" else sheet)
            if record:
                rows.append(record)

    rows.sort(key=lambda item: (item.get("academicYear", ""), item.get("company", ""), item.get("student", "")))
    return rows


def main():
    records = parse_workbook(SOURCE_XLSX)
    payload = {"records": records}
    OUTPUT_JSON.parent.mkdir(parents=True, exist_ok=True)
    OUTPUT_JSON.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"Generated: {OUTPUT_JSON}")
    print(f"Public records: {len(records)}")
    print("Sensitive fields excluded: phone, email, address, gender, DOB, district, etc.")


if __name__ == "__main__":
    main()
